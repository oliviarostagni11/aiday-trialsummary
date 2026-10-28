#!/usr/bin/env python3
"""Generate a coherent dummy dataset for the Cleo Diagnostics (CleoUSA) monthly trial summary.

Every output file mirrors the column layout and value formats of a real export:
  - seven Citrus CSV exports (participants, CRF and ePRO completeness, queries,
    inclusion, exclusion, post-surgical findings)
  - the LabConnect inventory report (.xlsx), built on the sample report as a style template

All files derive from one simulated participant journey, so the files agree with each other.
The export snapshot is taken on the first working day after the July 2026 month end.

Usage:
  .venv/bin/python scripts/generate_dummy_data.py [--template PATH] [--out DIR] [--seed N]
"""
import argparse
import copy
import csv
import random
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = Path("/Users/oliviarostagni/Desktop/AIbuildday Files/"
                        "CSP-01-001(LIND0004)_InventoryReport_18SEP2026.xlsx")
DEFAULT_OUT = ROOT / "data" / "dummy_inputs"

SEED = 20260731
EXPORT_AT = datetime(2026, 8, 3, 13, 56, 41, tzinfo=timezone.utc)
UTC = timezone.utc

SITE_LOOKUP = ROOT / "data" / "reference" / "site_lookup.csv"
SITE_MONTHLY = ROOT / "data" / "reference" / "site_monthly_enrollment.csv"
STUDY_START = date(2024, 8, 5)  # fallback activation date for a site with none recorded

# Extra participants created per enrolled participant, around the same month
NON_ENROLLED_RATES = (("prescreen", 0.30), ("screen_fail", 0.07), ("stale_consent", 0.04))

CRF_FORMS = [  # (encounter, form_title, form_variable) in the order Citrus exports them
    ("Enrolment Visit", "Informed Consent", "INFC"),
    ("Post-Surgery Data Collection", "Post-Surgical Findings", "SURG"),
    ("Post-Surgery Data Collection", "Study Completion", "COMPL"),
    ("Enrolment Visit", "Inclusion Criteria", "INCLC"),
    ("Enrolment Visit", "Exclusion Criteria", "EXCLC"),
    ("Enrolment Visit", "Enrollment", "ENROLL"),
    ("Enrolment Visit", "Study Blood Sample Collection & Processing", "BLOOD"),
    ("Enrolment Visit", "Health Information Collection", "HLTHI"),
    ("Enrolment Visit", "CA 125 Test, Pre-Surgery", "CATEST"),
    ("Enrolment Visit", "Demographics", "DEMOCRF"),
    ("Enrolment Visit", "Medical History", "MHQCRF"),
]
PRE_ENROL_FORMS = ["INCLC", "EXCLC", "DEMOCRF", "MHQCRF", "HLTHI", "CATEST"]
EPRO_FORMS = [("Demographics ePRO", "DEMOEPRO"), ("Medical History ePRO", "MHQEPRO")]

QUERY_QUESTIONS = {
    "Informed Consent": ["Has the participant provided informed consent to participate in this study?"],
    "Inclusion Criteria": ["#1 Are female patients greater than or equal to 18 years of age at time of consent",
                           "#2 Can provide written informed consent"],
    "Exclusion Criteria": ["#2 Any prior surgery resulting in removal of both ovaries",
                           "#3 Prior history of gynecological malignancy (within last 2 years)",
                           "#7 Pregnant currently or within the last 3 months based on participant self-report"],
    "Medical History": ["Do you currently take any prescription medications?",
                        "Do you have history of malignancy (cancer) in the last 5 years?",
                        "Have you or any of your family previously been diagnosed with Ovarian/Breast Cancer OR have Known Genetic Factors"],
    "Health Information Collection": ["Please upload the redacted radiological imaging report",
                                      "What is the role of the person making the overall impression of disease?",
                                      "Any additional comments?"],
    "CA 125 Test, Pre-Surgery": ["Please upload the redacted lab report of the CA 125 measurement",
                                 "What was the result of CA 125 measurement?",
                                 "Record the reason why the test result wasn't collected"],
    "Enrollment": ["Was the participant enrolled in the trial?"],
    "Study Blood Sample Collection & Processing": ["Upload the completed lab requisition form for the collected sample.",
                                                   "What was the date the sample was collected?",
                                                   "Record the reason why the sample wasn't collected or was collected with deviation."],
    "Adverse Events": ["Did the adverse event result in initial or prolonged hospitalization for the subject?",
                       "Did the adverse event result in disability or permanent damage?"],
    "Protocol Deviations": ["Provide a longer description of the deviation for context.", "Note date reported to IRB"],
    "Post-Surgical Findings": ["Please upload the redacted pathology report", "What date was the surgery performed?",
                               "What is the primary site of the malignancy?", "What was the malignancy stage?"],
    "Study Completion": ["What was the date and time participant completed all study activities?"],
}
FORM_TO_VAR = {f: v for _, f, v in CRF_FORMS}
STANDARD_QUERY_TEXT = [
    "Please confirm the date entered matches the source document.",
    "Uploaded file is not redacted. Please re-upload a redacted copy.",
    "Value is outside the expected range. Please verify.",
    "Date is before the consent date. Please correct.",
    "Please confirm the response against the medical record.",
]
DATA_MANAGERS = ["Priya Natarajan (Data Manager)", "Tom Becker (Data Manager)", "Lena Ortiz (Data Manager)"]
ASSIGNEES = {
    "clinical": ["Dana Whitfield (Investigator)", "Marcus Lee (Investigator)", "Hannah Cole (Trial Coordinator)"],
    "Adverse Events": ["Robert Kim (Medical Monitor)"],
    "Protocol Deviations": ["Sofia Reyes (Trial Monitor)"],
    "admin": ["Jordan Price (Admin)", "Alex Morgan (Admin)"],
}
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


# ---------- formatting helpers (match the export formats exactly) ----------
def ts(dt):
    return "NaT" if dt is None else dt.astimezone(UTC).strftime("%Y-%m-%d %H:%M:%S+00:00")


def pandas_timedelta(td):
    secs = int(td.total_seconds())
    micro = td.microseconds
    days, rem = divmod(secs, 86400)
    h, rem = divmod(rem, 3600)
    m, s = divmod(rem, 60)
    return f"{days} days {h:02d}:{m:02d}:{s:02d}.{micro:06d}"


def lab_date(d):  # "22 Nov 2024"
    return f"{d.day:02d} {MONTHS[d.month - 1]} {d.year}"


def dob(pid):
    """Fake date of birth, aged 18 to 80 at 2026. Seeded per participant so it does not shift the main random stream."""
    r = random.Random(f"dob-{pid}")
    return date(2026 - r.randint(18, 80), r.randint(1, 12), r.randint(1, 28))


def ship_date(d):  # "6-Jul-2026 0:00:00"
    return f"{d.day}-{MONTHS[d.month - 1]}-{d.year} 0:00:00"


def comment_date(d):  # "27JUN2025"
    return f"{d.day:02d}{MONTHS[d.month - 1].upper()}{d.year}"


class Gen:
    def __init__(self, seed):
        self.r = random.Random(seed)

    # random helpers
    def chance(self, p):
        return self.r.random() < p

    def poisson(self, lam):
        n, p, limit = 0, 1.0, pow(2.718281828459045, -lam)
        while True:
            p *= self.r.random()
            if p <= limit:
                return n
            n += 1

    def clock(self, d, start_h=8, end_h=17, seconds=False):
        t = time(self.r.randint(start_h, end_h - 1), self.r.randrange(0, 60, 1 if seconds else 5),
                 self.r.randint(0, 59) if seconds else 0)
        return datetime.combine(d, t, tzinfo=UTC)

    def jitter(self, dt, max_minutes=30):
        return dt + timedelta(seconds=self.r.randint(5, max_minutes * 60), microseconds=self.r.randint(0, 999999))

    def citrus_path(self):
        alphabet = "23456789ABCDEFGHJKLMNPQRSTUVWXYZabcdefghijkmnopqrstuvwxyz"
        key = "".join(self.r.choice(alphabet) for _ in range(22))
        return f"permanent/study/k3BMbCaZduGEqwsyXLBemh/dashboard/{key}/{uuid.UUID(int=self.r.getrandbits(128), version=4)}"


# ---------- participant journeys ----------
def load_sites():
    """Sites and their recorded enrollments per month (from scripts/import_site_list.py)."""
    with open(SITE_LOOKUP, encoding="utf-8-sig") as f:
        sites = list(csv.DictReader(f))
    monthly = defaultdict(dict)
    with open(SITE_MONTHLY, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            monthly[r["lab_site_code"]][r["month"]] = int(r["enrolled"])
    return sites, monthly


def simulate_participants(g, sites, monthly):
    """Enrollments per site and month match the site performance workbook exactly.
    Pre-screen-only, screen-failure and stalled-consent participants are added around them."""
    people = []
    for s in sites:
        act = date.fromisoformat(s["activation_date"]) if s["activation_date"] else STUDY_START
        for month, n in sorted(monthly[s["lab_site_code"]].items()):
            y, m = (int(x) for x in month.split("-"))
            first = date(y, m, 1)
            last = date(y + (m == 12), m % 12 + 1, 1) - timedelta(days=1)
            if first <= act <= last:
                first = act
            days = [first + timedelta(days=i) for i in range((last - first).days + 1)]
            days = [d for d in days if d.weekday() < 5] or days
            kinds = ["enrolled"] * n
            for kind, rate in NON_ENROLLED_RATES:
                kinds += [kind] * sum(g.chance(rate) for _ in range(n))
            for kind in kinds:
                p = {"site_id": s["site_id"], "site": s["site"]}
                if journey(g, p, kind, g.r.choice(days)):
                    people.append(p)
    people.sort(key=lambda p: p["prescreen"])
    for i, p in enumerate(people, start=1):
        p["pid"] = f"USA-{i:03d}"
    return people


def journey(g, p, kind, d):
    """Set each participant's milestones. Anything after EXPORT_AT has not happened yet.
    Returns False when the participant does not exist yet at the export time."""
    E = EXPORT_AT
    p.update(consent=None, enrolled=None, completed=None, sf_at=None, sf_reason="", wd_at=None, wd_reason="",
             surgery=None, surgery_entered=None, path_upload=None, path_never=False, no_sample=False,
             diagnosis=None, status="PRE-SCREENING PASSED")
    if kind == "prescreen":
        p["prescreen"] = g.clock(d, seconds=True)
        return p["prescreen"] <= E
    if kind == "enrolled":
        enrolled = g.clock(d)
        consent = enrolled - (timedelta(minutes=g.r.choice([20, 35, 50])) if g.chance(0.55)
                              else timedelta(days=g.r.randint(1, 14), minutes=g.r.randint(0, 300)))
    else:
        enrolled, consent = None, g.clock(d)
    p["prescreen"] = consent - timedelta(days=g.r.choice([0, 0, 1, 3, 7, 14, 21]), hours=g.r.randint(1, 5),
                                         seconds=g.r.randint(0, 59))
    if p["prescreen"] > E:
        return False
    if consent > E:
        return True
    p["consent"], p["status"] = consent, "CONSENT"
    if kind == "screen_fail":
        sf = consent + timedelta(days=g.r.randint(0, 14), hours=g.r.randint(0, 6))
        if sf <= E:
            p["sf_at"], p["status"] = sf, "SCREEN FAILURE"
            p["sf_reason"] = g.r.choice(["Failure To Meet Inclusion/Exclusion Criteria"] * 3 + ["Consent Declined"])
        return True
    if kind == "stale_consent" or enrolled > E:
        return True
    p["enrolled"], p["status"] = enrolled, "ENROLLED"
    surgery_planned = enrolled + timedelta(days=round(g.r.triangular(2, 45, 14)))
    if g.chance(0.05):
        wd = enrolled + timedelta(days=g.r.randint(1, 40))
        if wd <= E and wd < surgery_planned:
            p["wd_at"], p["status"] = g.clock(wd.date()), "WITHDRAWAL"
            p["wd_reason"] = g.r.choice([
                "Unable to Obtain Study Blood Sample",
                "Other:: Cyst was shown to be resolved upon follow up imaging.",
                "Other:: Pt's cyst was resolved upon follow up imaging; no longer indicated for surgery.",
                "Adverse Event"])
            p["no_sample"] = p["wd_reason"] == "Unable to Obtain Study Blood Sample"
            return True
    if surgery_planned.date() > E.date():
        return True
    p["surgery"] = surgery_planned.date()
    lag = g.r.randint(0, 21) if g.chance(0.9) else g.r.randint(22, 60)
    entered = g.clock(p["surgery"] + timedelta(days=lag), seconds=True)
    if entered > E:
        return True
    p["surgery_entered"] = entered
    p["diagnosis"] = g.r.choices(["Benign [0]", "Borderline/LMP [1]", "Malignant [2]"], [60, 10, 30])[0]
    if g.chance(0.07):
        p["path_never"] = True
        return True
    upload = g.clock(p["surgery"] + timedelta(days=round(g.r.triangular(7, 50, 18))), seconds=True)
    upload = max(upload, entered + timedelta(minutes=g.r.randint(2, 40)))
    if upload > E:
        return True
    p["path_upload"] = upload
    done = g.clock((upload + timedelta(days=g.r.randint(0, 10))).date())
    done = max(done, upload + timedelta(minutes=10))
    if done <= E:
        p["completed"], p["status"] = done, "COMPLETED"
    return True


# ---------- queries ----------
def simulate_queries(g, people):
    rows, open_forms = [], set()
    E = EXPORT_AT
    for p in people:
        if not p["consent"]:
            continue
        lam = 0.35 + (0.5 if p["enrolled"] else 0) + (0.35 if p["surgery_entered"] else 0)
        end = min(E, (p["completed"] or p["wd_at"] or p["sf_at"] or E) + timedelta(days=45))
        forms = ["Informed Consent", "Inclusion Criteria", "Exclusion Criteria", "Medical History",
                 "Health Information Collection", "CA 125 Test, Pre-Surgery"]
        if p["enrolled"]:
            forms += ["Enrollment", "Study Blood Sample Collection & Processing", "Adverse Events",
                      "Protocol Deviations", "Health Information Collection", "CA 125 Test, Pre-Surgery"]
        if p["surgery_entered"]:
            forms += ["Post-Surgical Findings", "Post-Surgical Findings", "Study Completion"]
        for _ in range(g.poisson(lam)):
            start = p["consent"] + timedelta(days=1)
            if start >= end:
                continue
            created = start + timedelta(seconds=g.r.randint(0, int((end - start).total_seconds())))
            created = created.replace(microsecond=0)
            form = g.r.choice(forms)
            qtype = "Standard" if g.chance(0.22) else "Missing Data"
            stale = g.chance(0.12)
            close_at = None if stale else created + timedelta(days=g.r.expovariate(1 / 12), seconds=g.r.randint(60, 7200))
            if close_at and close_at <= E:
                status, updated = "Closed", close_at.replace(microsecond=g.r.randint(0, 999999))
            elif g.chance(0.3):
                status = "Reviewing"
                updated = created + timedelta(seconds=g.r.randint(60, max(61, int((E - created).total_seconds() * 0.5))))
            else:
                status, updated = "Open", created
            age_end = updated if status == "Closed" else E.replace(microsecond=g.r.randint(0, 999999))
            if form in ASSIGNEES:
                assignee = g.r.choice(ASSIGNEES[form])
            else:
                assignee = g.r.choice(ASSIGNEES["clinical"] * 3 + ASSIGNEES["admin"])
            question = g.r.choice(QUERY_QUESTIONS[form])
            rows.append({
                "Participant ID": p["pid"], "Site ID": p["site_id"], "Participant Status": p["status"],
                "Site": p["site"], "Created At": ts(created), "Age": pandas_timedelta(age_end - created),
                "Updated at": ts(updated),
                "Query content": g.r.choice(STANDARD_QUERY_TEXT) if qtype == "Standard" else "",
                "Type": qtype, "Status": status, "Form Title": form, "Input question": question,
                "Input string value": (g.r.choice(["1", "0", "n/a", "2026-01-14"]) if qtype == "Standard" else ""),
                "Reporter name": g.r.choice(DATA_MANAGERS) if qtype == "Standard" else "  ()",
                "Assignee name": assignee,
                "_created": created,
            })
            if status == "Open" and form in FORM_TO_VAR:
                open_forms.add((p["pid"], FORM_TO_VAR[form]))
    rows.sort(key=lambda q: (q["Participant ID"], q["_created"]))
    return rows, open_forms


# ---------- CRF completeness ----------
def crf_statuses(g, p, open_forms):
    E = EXPORT_AT
    s = {v: "Unavailable" for _, _, v in CRF_FORMS}
    if p["status"] == "PRE-SCREENING PASSED":
        s["INFC"] = "Available"
    elif p["status"] == "SCREEN FAILURE":
        s.update(INFC="Complete", INCLC="Complete", EXCLC="Complete")
        for v in ("DEMOCRF", "MHQCRF"):
            s[v] = "Complete" if g.chance(0.5) else "Available"
    elif p["status"] == "CONSENT":
        s["INFC"] = "Complete"
        for v in PRE_ENROL_FORMS:
            s[v] = "Complete" if g.chance(0.6) else "Available"
        s["ENROLL"], s["BLOOD"] = "Available", "Available"
    else:  # ENROLLED, COMPLETED, WITHDRAWAL
        recent = (E - p["enrolled"]).days <= 30
        s["INFC"] = "Complete"
        for v in PRE_ENROL_FORMS + ["ENROLL"]:
            s[v] = "Complete" if g.chance(0.78 if recent else 0.96) else "Available"
        s["BLOOD"] = "Skipped" if p["no_sample"] else ("Complete" if g.chance(0.85 if recent else 0.97) else "Available")
        if p["status"] == "WITHDRAWAL":
            s["SURG"], s["COMPL"] = "Skipped", "Complete"
        else:
            s["SURG"] = "Complete" if p["path_upload"] else "Available"
            s["COMPL"] = "Complete" if p["completed"] else ("Available" if p["path_upload"] else "Unavailable")
        if p["completed"] and p["completed"] < datetime(2024, 12, 31, tzinfo=UTC):
            s = {v: ("Locked" if st == "Complete" else st) for v, st in s.items()}
    for v in s:
        if (p["pid"], v) in open_forms and s[v] != "Unavailable":
            s[v] = "Open Query"
    return s


# ---------- inclusion / exclusion ----------
INCL_COLS = ["#1 Are female patients greater than or equal to 18 years of age at time of consent (IC#1)",
             "#2 Can provide written informed consent (IC#2)",
             "#3 Have an ovarian cyst or adnexal mass (simple, complex, or solid ovarian/pelvic mass) identified through radiological imaging examination that requires surgery, but has not yet undergone this surgery (IC#3)"]
EXCL_COLS = ["#1 Any prior confirmed diagnosis of, or treatment for, ovarian cancer (EC#1)",
             "#2 Any prior surgery resulting in removal of both ovaries (EC#2)",
             "#3 Prior history of gynecological malignancy (within last 2 years) (EC#3)",
             "#4 Prior history of melanoma (within last 2 years) (EC#4)",
             "#5 Prior treatment (within 12 months of sample collection) with chemotherapy, radiotherapy or immunotherapy (EC#5)",
             "#6 Immune-compromised. Definition: those whose weakened immune system may render them more susceptible to infection, illness, or complications. These may include patients with immune-suppressive conditions (e.g. HIV/AIDS), patients undergoing immuno-, chemo- or radio-therapy (within the last 12 months) for any reason, receiving immune suppressive therapies (e.g. transplant recipients, individuals with known autoimmune disease), or with a primary immunodeficiency disorder (e.g. Lupus). (EC#6)",
             "#7 Pregnant currently or within the last 3 months based on participant self-report (EC#7)"]
FORM_META_COLS = ["Participant ID", "Site ID", "Participant Status", "Site", "Encounter", "Form Title", "Form Status",
                  "Available At", "Completed At", "Last edited by", "Last edited at", "Skipped At", "Skipped Reason"]


def editor(p):
    return f"CleoUSA: crc.{p['site_id'][4:].lower()}@example.org (None)"


def ie_rows(g, p, crf, which):
    var, cols = ("INCLC", INCL_COLS) if which == "incl" else ("EXCLC", EXCL_COLS)
    st = crf[var]
    if st == "Unavailable" or (st == "Available" and not g.chance(0.4)):
        return None
    complete = st in ("Complete", "Locked")
    avail = g.jitter(p["consent"], 2)
    done = g.jitter(p["consent"], 25) if complete else None
    last = done or g.jitter(p["consent"], 40)
    fail_ie = p["status"] == "SCREEN FAILURE" and p["sf_reason"].startswith("Failure")
    if which == "incl":
        answers = ["Yes [1]"] * 3
        if fail_ie and g.chance(0.5):
            p["_ie_fail_side"] = "incl"
            answers[2] = "No [0]"
    else:
        answers = ["No [0]"] * 7
        if g.chance(0.08):
            answers[6] = "N/A [99]"
        if fail_ie and p.get("_ie_fail_side") != "incl":
            answers[g.r.choice([0, 1, 2, 4])] = "Yes [1]"
    if not complete:
        for i in g.r.sample(range(len(answers)), k=min(len(answers), g.r.randint(1, 3))):
            answers[i] = "Suspected missing"
    row = {"Participant ID": p["pid"], "Site ID": p["site_id"], "Participant Status": p["status"],
           "Site": p["site"], "Encounter": "Enrolment Visit",
           "Form Title": "Inclusion Criteria" if which == "incl" else "Exclusion Criteria",
           "Form Status": "COMPLETE" if complete else "IN_PROGRESS",
           "Available At": ts(avail), "Completed At": ts(done), "Last edited by": editor(p),
           "Last edited at": ts(last), "Skipped At": "", "Skipped Reason": ""}
    row.update(dict(zip(cols, answers)))
    return row


# ---------- post-surgical findings ----------
PSF_COLS = ["What is the role of the person who performed the surgery? (Role of person making pre-surgical impression)",
            "What date was the surgery performed? (Surgery Date)",
            "What was the post-surgical diagnosis? (Cancer Diagnosis)",
            "What is the primary site of the malignancy? (Malignancy Site)",
            "Detail pathology sub type (Pathology Type/Sub-Type)",
            "What was the malignancy stage? (Malignancy Stage)",
            "What was the malignancy sub-stage? (Malignancy Sub-Stage)",
            "What was the malignancy grade? (Malignancy Grade)",
            "Was STIC-Serous Tubal Intra-epithelial carcinoma present? (STIC)",
            "Is a P53 signature present? (P53 signature)",
            "Please upload the redacted pathology report (Pathology Report)"]
STAGES = [("Stage I: Cancer confined to one or both ovaries [0]",
           ["IA: Tumor limited to 1 ovary or fallopian tube, capsule intact, no tumor on surface, negative washings [0]",
            "IC1: Tumor limited to 1 or both ovaries or fallopian tubes with Surgical spill [2]",
            "IC2: Tumor limited to 1 or both ovaries or fallopian tubes with Capsule rupture before surgery or tumor on ovarian surface [3]"]),
          ("Stage II: Cancer involves one or both ovaries with pelvic extension [1]", ["No substage collected [88]"]),
          ("Stage III: Cancer involves one or both ovaries with microscopically confirmed peritoneal metastasis outside the pelvis and/or regional lymph node metastasis [2]",
           ["IIIA1(ii): Positive retroperitoneal lymph nodes only Metastasis > 10 mm [8]", "No substage collected [88]"]),
          ("Primary peritoneal [4]", ["No substage collected [88]"])]
GRADES = ["Grade 1 - low grade (well differentiated) [0]",
          "Grade 2 or Grade 3 - high grade (moderately or poorly differentiated) [1]", "Ungraded [2]"]


def psf_row(g, p):
    if not p["surgery_entered"]:
        return None
    complete = p["path_upload"] is not None
    vals = ["Gynecological Oncologist (GO) [1]", p["surgery"].isoformat()]
    if not complete:
        vals += ["Suspected missing"] * 8 + ["Suspected missing"]
    else:
        dx = p["diagnosis"]
        if dx == "Malignant [2]":
            stage, subs = g.r.choices(STAGES, [35, 15, 40, 10])[0]
            vals += [dx, g.r.choices(["Ovary", "Fallopian tube", "Other"], [70, 22, 8])[0],
                     g.r.choice(["High-grade serous carcinoma", "Low-grade serous carcinoma", "Endometrioid carcinoma",
                                 "Clear cell carcinoma", "Mucinous carcinoma"]),
                     stage, g.r.choice(subs), g.r.choices(GRADES, [30, 60, 10])[0],
                     "Yes [1]" if g.chance(0.25) else "No [0]", "Yes [1]" if g.chance(0.6) else "No [0]"]
        elif dx == "Borderline/LMP [1]":
            vals += [dx, "Ovary", g.r.choice(["Serous borderline tumor", "Mucinous borderline tumor"]),
                     "NOT ASKED", "NOT ASKED", "NOT ASKED", "NOT ASKED", "Yes [1]" if g.chance(0.2) else "No [0]"]
        else:
            vals += [dx] + ["NOT ASKED"] * 6 + ["Yes [1]" if g.chance(0.15) else "No [0]"]
        vals.append(g.citrus_path())
    row = {"Participant ID": p["pid"], "Site ID": p["site_id"], "Participant Status": p["status"],
           "Site": p["site"], "Encounter": "Post-Surgery Data Collection", "Form Title": "Post-Surgical Findings",
           "Form Status": "COMPLETE" if complete else "IN_PROGRESS",
           "Available At": ts(g.jitter(p["enrolled"], 5)), "Completed At": ts(p["path_upload"]),
           "Last edited by": editor(p), "Last edited at": ts(p["path_upload"] or p["surgery_entered"]),
           "Skipped At": "", "Skipped Reason": ""}
    row.update(dict(zip(PSF_COLS, vals)))
    return row


# ---------- ePRO ----------
def epro_rows(g, p):
    if not p["consent"]:
        return []
    out = []
    for title, var in EPRO_FORMS:
        c = dict(num_available=0, num_expired=0, num_completed=0, num_skipped=0, supp_num_overdue=0)
        age = (EXPORT_AT - p["consent"]).days
        if age <= 3:
            c["num_available"] = 1
        else:
            pick = g.r.choices(["completed", "expired", "skipped", "available"], [80, 13, 3, 4])[0]
            c["num_" + pick] = 1
            if pick == "available" and age > 7:
                c["supp_num_overdue"] = 1
        out.append({"participant_id": p["pid"], "site_id": p["site_id"], "site": p["site"],
                    "encounter": "Enrolment Visit", "form_title": title, "form_variable": var, **c})
    return out


# ---------- inventory report ----------
def simulate_inventory(g, people):
    """One requisition per enrolled participant who gave blood. DX1 and DX2 share a collection date."""
    E = EXPORT_AT.date()
    samples, req = [], 48390700
    scenario_counts = Counter()
    lost = {p["pid"] for p in g.r.sample([p for p in people if p["status"] == "ENROLLED"
                                          and (EXPORT_AT - p["enrolled"]).days > 30], k=2)}
    for p in people:
        if not p["enrolled"] or p["no_sample"]:
            continue
        if p["pid"] in lost:
            scenario_counts["no sample registered"] += 1
            continue
        req += g.r.randint(1, 9)
        coll = p["enrolled"].date()
        coll_time = (p["enrolled"] + timedelta(minutes=g.r.choice([10, 15, 20, 30]))).strftime("%H:%M:00")
        scen = g.r.choices(["normal", "dx1_late", "dx2_late", "same_day", "dx2_missing", "dx1_missing"],
                           [74, 10, 8, 4, 3, 1])[0]
        dx1 = coll + timedelta(days=g.r.randint(1, 10))
        dx2 = coll + timedelta(days=g.r.randint(12, 26))
        if scen == "dx1_late":
            dx1 = coll + timedelta(days=g.r.randint(15, 45))
            dx2 = max(dx2, dx1 + timedelta(days=g.r.randint(2, 10)))
        elif scen == "dx2_late":
            dx2 = coll + timedelta(days=g.r.randint(29, 60))
        elif scen == "same_day":
            dx2 = dx1
        if dx2 == dx1 and scen != "same_day":
            dx2 += timedelta(days=1)
        regs = {"CLEODX1": None if scen == "dx1_missing" else dx1, "CLEODX2": None if scen == "dx2_missing" else dx2}
        scenario_counts[scen] += 1
        destroyed = g.chance(0.01)
        for code, reg in regs.items():
            if reg is None or reg > E:
                continue
            n = "1" if code == "CLEODX1" else "2"
            comment, open_q = "", "No"
            if destroyed:
                status, facility, ship = "Destroyed", "LabConnect JC Biohazard", ""
                comment = f"Sample hemolyzed; destroyed per sponsor instruction. CPH {comment_date(reg + timedelta(days=3))}."
            else:
                status, facility, ship = "In Inventory", "BRS Biobank", ""
                if reg <= date(2025, 11, 28):
                    ship = ship_date(date(2025, 12, 15) if code == "CLEODX1" else date(2025, 12, 17))
                elif reg <= date(2026, 2, 20) and g.chance(0.05):
                    ship = ship_date(date(2026, 3, 4))
                elif date(2026, 2, 20) < reg <= date(2026, 6, 19) and g.chance(0.6):
                    ship = ship_date(date(2026, 7, 6))
                if ship:
                    status, facility = "Transferred", "Cleo Diagnostics"
                if g.chance(0.07):
                    comment = g.r.choice([
                        f"SID on tube is missing the leading 0.RCS,{comment_date(reg)}.",
                        f"SID on tube is illegible.RCS,{comment_date(reg)}.",
                        f"No SID on tube.RCS,{comment_date(reg)}."])
                    if g.chance(0.25):
                        open_q = "Yes"
                        comment = f"Query submitted for missing SID. {comment_date(reg)} LBK"
            samples.append({
                "Sponsor": "Cleo Diagnostics Ltd", "Protocol": "CSP-01-001", "LC Name": "LIND0004",
                "Requisition Number": str(req), "Accession Number": f"{req}10{n}01",
                "LabMatrix Barcode": f"{req}-{code}", "Site": p["site_id"][4:], "Subject ID": p["pid"],
                "Date of Birth": lab_date(dob(p["pid"])), "Visit Type": "Screening", "Date of Collection": lab_date(coll),
                "Time of Collection": coll_time, "Date Registered": lab_date(reg), "Test Code": code,
                "Test Name": f"{code} Storage Sample {n}", "LabREACH Comments": "", "Specimen Comments": comment,
                "Current Status": status, "Shipment Date": ship, "At Facility": facility,
                "Is there an open query?": open_q,
            })
    samples.sort(key=lambda s: (s["Test Code"], s["Requisition Number"]))
    return samples, scenario_counts, lost


def write_inventory(samples, template, out_path):
    wb = openpyxl.load_workbook(template)
    ws = wb.active
    header = [c.value for c in ws[4]]
    body_style = [copy.copy(c._style) for c in ws[5]]
    ws.delete_rows(5, ws.max_row)
    ws["A3"] = f"Date: {EXPORT_AT.day:02d}{MONTHS[EXPORT_AT.month - 1].upper()}{EXPORT_AT.year}"
    for i, s in enumerate(samples, start=5):
        for j, h in enumerate(header, start=1):
            cell = ws.cell(row=i, column=j, value=(s[h] if s[h] != "" else None))
            cell._style = copy.copy(body_style[j - 1])
    ws.freeze_panes = "A5"
    wb.save(out_path)


def write_csv(path, rows, cols):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, lineterminator="\n", extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--seed", type=int, default=SEED)
    a = ap.parse_args()
    a.out.mkdir(parents=True, exist_ok=True)
    g = Gen(a.seed)
    stamp = lambda minute: f"CleoUSA-{EXPORT_AT:%Y-%m-%d}-13-{minute}-UTC.csv"

    people = simulate_participants(g, *load_sites())
    queries, open_forms = simulate_queries(g, people)

    write_csv(a.out / f"PARTICIPANTS_OVERVIEW-{stamp(57)}", [{
        "Participant ID": p["pid"], "Participant Status": p["status"], "Site ID": p["site_id"], "Site": p["site"],
        "Screen failure at": ts(p["sf_at"]), "Screen failure reason": p["sf_reason"],
        "Withdrawal at": ts(p["wd_at"]), "Withdrawal reason": p["wd_reason"], "Consent at": ts(p["consent"]),
        "Enrolled at": ts(p["enrolled"]), "Completed at": ts(p["completed"])} for p in people],
        ["Participant ID", "Participant Status", "Site ID", "Site", "Screen failure at", "Screen failure reason",
         "Withdrawal at", "Withdrawal reason", "Consent at", "Enrolled at", "Completed at"])

    crf_rows, incl, excl, psf, epro = [], [], [], [], []
    for p in people:
        crf = crf_statuses(g, p, open_forms)
        for enc, title, var in CRF_FORMS:
            crf_rows.append({"participant_id": p["pid"], "site_id": p["site_id"], "site": p["site"],
                             "encounter": enc, "form_title": title, "form_variable": var, "status": crf[var]})
        for which, bucket in (("incl", incl), ("excl", excl)):
            row = ie_rows(g, p, crf, which)
            if row:
                bucket.append(row)
        row = psf_row(g, p)
        if row:
            psf.append(row)
        epro += epro_rows(g, p)

    write_csv(a.out / f"DATA_COMPLETENESS__CRF-{stamp(59)}", crf_rows,
              ["participant_id", "site_id", "site", "encounter", "form_title", "form_variable", "status"])
    write_csv(a.out / f"DATA_COMPLETENESS__EPRO-{stamp(59)}", epro,
              ["participant_id", "site_id", "site", "encounter", "form_title", "form_variable", "num_available",
               "num_expired", "num_completed", "num_skipped", "supp_num_overdue"])
    write_csv(a.out / f"DATA_QUERY-{stamp(58)}", queries,
              ["Participant ID", "Site ID", "Participant Status", "Site", "Created At", "Age", "Updated at",
               "Query content", "Type", "Status", "Form Title", "Input question", "Input string value",
               "Reporter name", "Assignee name"])
    write_csv(a.out / f"Inclusion_Criteria-{stamp(58)}", incl, FORM_META_COLS + INCL_COLS)
    write_csv(a.out / f"Exclusion_Criteria-{stamp(58)}", excl, FORM_META_COLS + EXCL_COLS)
    write_csv(a.out / f"PostSurgical_Findings-{stamp(56)}", psf, FORM_META_COLS + PSF_COLS)

    samples, scen, lost = simulate_inventory(g, people)
    inv_name = f"CSP-01-001(LIND0004)_InventoryReport_{EXPORT_AT.day:02d}{MONTHS[EXPORT_AT.month - 1].upper()}{EXPORT_AT.year}.xlsx"
    write_inventory(samples, a.template, a.out / inv_name)
    # CSV copy of the inventory data rows (header row onward) for quick viewing
    write_csv(a.out / inv_name.replace('.xlsx', '.csv'), samples, list(samples[0].keys()))

    # ---- console summary for a quick sanity check ----
    in_jul = lambda dt: dt is not None and dt.year == 2026 and dt.month == 7
    print(f"Export snapshot: {ts(EXPORT_AT)}  ->  {a.out}")
    print("Participants:", len(people), dict(Counter(p['status'] for p in people)))
    print("July 2026  consents:", sum(in_jul(p['consent']) for p in people),
          " enrolments:", sum(in_jul(p['enrolled']) for p in people),
          " completions:", sum(in_jul(p['completed']) for p in people),
          " surgeries:", sum(p['surgery'] is not None and p['surgery'].year == 2026 and p['surgery'].month == 7 for p in people),
          " pathology uploads:", sum(in_jul(p['path_upload']) for p in people),
          " queries created:", sum(in_jul(q['_created']) for q in queries))
    print("Queries:", len(queries), dict(Counter(q['Status'] for q in queries)), dict(Counter(q['Type'] for q in queries)))
    print("CRF rows:", len(crf_rows), dict(Counter(r['status'] for r in crf_rows)))
    print("ePRO rows:", len(epro), " Inclusion rows:", len(incl), " Exclusion rows:", len(excl), " PSF rows:", len(psf))
    print("Inventory samples:", len(samples), " pair scenarios:", dict(scen), " lost:", sorted(lost))


if __name__ == "__main__":
    main()
