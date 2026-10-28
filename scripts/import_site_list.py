#!/usr/bin/env python3
"""Import the CLEO-001 site performance workbook into the reference files the report and generator use.

Writes to data/reference/:
  site_lookup.csv              one row per site: IDs, name, type, academic status, activation, target, status
  site_monthly_enrollment.csv  enrollments per site per month, as recorded in the workbook

Academic status comes from site_academic_status.csv (edit that file to change a site).

Usage:
  .venv/bin/python scripts/import_site_list.py ["/path/to/CLEO-001 Overall Site Performance by Month (2025).xlsx"]
"""
import csv
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "data" / "reference"
DEFAULT_SRC = Path("/Users/oliviarostagni/Desktop/AIbuildday Files/CLEO-001 Overall Site Performance by Month (2025).xlsx")

# Sites listed without a code in the workbook: code taken from the Citrus site ID.
CODE_FOR_NAME = {"ES Clinical Research Group": "DHW"}
TYPE_LABELS = {"Gyn/Onc.": "Gyn/Onc", "General Gyn": "General Gyn", "Gen Onc": "General Onc"}
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def month_key(header):
    m = re.match(r"([A-Za-z]+)\s+(\d{4})", str(header or ""))
    if not m:
        return None
    return f"{m.group(2)}-{MONTHS[m.group(1)[:3].lower()]:02d}"


def count(v):
    """Workbook cells hold numbers, 'NA' (not active), 'Activated-0', 'Term' or '0- Terminated 03 July'."""
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v or "").strip()
    m = re.match(r"(?:Activated-)?(\d+)", s)
    return int(m.group(1)) if m else None


def target(v):
    if v is None:
        return ""
    return int(v) if isinstance(v, float) and v.is_integer() else v


def main():
    src = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SRC
    ws = openpyxl.load_workbook(src, data_only=True).worksheets[0]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header)}
    month_cols = [(i, month_key(h)) for i, h in enumerate(header) if month_key(h)]
    academic = {r["lab_site_code"]: r for r in csv.DictReader(open(REF / "site_academic_status.csv", encoding="utf-8"))}

    sites, monthly = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        raw = (r[col["Site Name"]] or "").strip()
        if not raw or raw == "Total Enrollment":
            continue
        m = re.match(r"([A-Z]{3})\s*-\s*(.+)", raw)
        code, name = (m.group(1), m.group(2).strip()) if m else (CODE_FOR_NAME[raw], raw)
        name = re.sub(r"\s+", " ", name).replace("’", "'")
        notes = str(r[col["Notes"]] or "")
        terminated = "terminat" in notes.lower()
        term_date = ""
        for i, mk in month_cols:
            cell = str(r[i] or "")
            t = re.search(r"Terminated (\d{1,2}) ([A-Za-z]+)", cell)
            if t:
                term_date = date(int(mk[:4]), MONTHS[t.group(2)[:3].lower()], int(t.group(1))).isoformat()
        act = r[col["Activation Date"]]
        acad = academic.get(code)
        sites.append({
            "site_id": f"CDX-{code}", "lab_site_code": code, "site": name,
            "site_type": TYPE_LABELS.get((r[col["Type"]] or "").strip(), (r[col["Type"]] or "").strip()),
            "academic_status": acad["academic_status"] if acad else "Unknown",
            "activation_date": act.date().isoformat() if isinstance(act, datetime) else "",
            "monthly_target": target(r[col["Monthly Target"]]),
            "tracker_flag": r[col["Site Deactivation"]] or "",
            "terminated": "Yes" if terminated else "No", "termination_date": term_date,
        })
        for i, mk in month_cols:
            n = count(r[i])
            if n:
                monthly.append({"lab_site_code": code, "month": mk, "enrolled": n})

    fields = list(sites[0].keys())
    with open(REF / "site_lookup.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, lineterminator="\n")
        w.writeheader()
        w.writerows(sites)
    with open(REF / "site_monthly_enrollment.csv", "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["lab_site_code", "month", "enrolled"], lineterminator="\n")
        w.writeheader()
        w.writerows(monthly)
    print(f"{len(sites)} sites, {sum(m['enrolled'] for m in monthly)} enrollments across {len({m['month'] for m in monthly})} months")
    missing = [s["lab_site_code"] for s in sites if s["academic_status"] == "Unknown"]
    if missing:
        print("No academic status for:", ", ".join(missing))


if __name__ == "__main__":
    main()
