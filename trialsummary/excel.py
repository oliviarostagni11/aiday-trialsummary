"""Write the client workbook: a summary tab, one tab per section and a definitions tab."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import charts, narrative
from .audience import AUDIENCES, sections_for, tables_for

FONT = "Arial"
NAVY, BLUE, BLUE_LIGHT = "184F95", "1C5CAB", "86B6EF"
INK, INK_2, MUTED = "0B0B0B", "52514E", "898781"
ZEBRA, TOTAL_FILL, RULE = "F4F7FB", "E6EEF8", "C3C2B7"
THIN = Side(style="thin", color="E1E0D9")
TOP = Side(style="thin", color=RULE)

TAB_NAMES = {"enrollment": "Enrollment", "completeness": "Data completeness", "post_surgical": "Post-surgical",
             "blood": "Blood samples", "queries": "Data queries"}
NUM_FMT = {"int": "#,##0", "pct": '[=0]0%;[<0.005]"<1%";0%', "dec": "0.0", "rate": "0.00", "date": "d mmm yyyy"}


def _font(**kw):
    kw.setdefault("name", FONT)
    kw.setdefault("size", 10)
    kw.setdefault("color", INK)
    return Font(**kw)


def _value(col, v):
    if v is None:
        return "–", None
    if col.kind == "range":
        lo, hi = v
        return f"{lo:.0f}–{hi:.0f}", None
    if col.kind == "int" and isinstance(v, float):
        return v, "0.00"
    return v, NUM_FMT.get(col.kind)


def _sheet_header(ws, title, subtitle, width_cols=10):
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = _font(size=16, bold=True, color=NAVY)
    ws["A2"] = subtitle
    ws["A2"].font = _font(size=9, color=INK_2)
    ws.row_dimensions[1].height = 24
    return 4


def _write_table(ws, t, row):
    ncol = len(t.cols)
    ws.cell(row=row, column=1, value=t.title).font = _font(size=12, bold=True, color=NAVY)
    row += 1
    if t.note:
        c = ws.cell(row=row, column=1, value=t.note)
        c.font = _font(size=8.5, italic=True, color=INK_2)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(ncol, 6))
        ws.row_dimensions[row].height = 13 * max(1, len(t.note) // 150 + 1)
        row += 1
    for j, col in enumerate(t.cols, start=1):
        c = ws.cell(row=row, column=j, value=col.header)
        c.font = _font(size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="left" if j == 1 else "center", vertical="center", wrap_text=True)
    lines = 1
    for j, col in enumerate(t.cols, start=1):
        width = ws.column_dimensions[get_column_letter(j)].width or 12
        lines = max(lines, -(-len(col.header) // max(int(width * 0.95), 1)))
    ws.row_dimensions[row].height = 13 * lines + 6
    header_row = row
    if not t.rows:
        row += 1
        ws.cell(row=row, column=1, value="None").font = _font(size=9, italic=True, color=INK_2)
        return row + 3
    for i, r in enumerate(t.rows):
        row += 1
        total = i in t.total_rows
        for j, (col, v) in enumerate(zip(t.cols, r), start=1):
            val, fmt = _value(col, v)
            c = ws.cell(row=row, column=j, value=val)
            c.font = _font(size=9.5, bold=total)
            if fmt:
                c.number_format = fmt
            c.alignment = Alignment(horizontal="left" if col.kind == "text" else "right" if val != "–" else "center",
                                    vertical="center")
            c.border = Border(bottom=THIN, top=TOP if total else None)
            if total:
                c.fill = PatternFill("solid", fgColor=TOTAL_FILL)
            elif i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
    body = [header_row + 1 + i for i in range(len(t.rows)) if i not in t.total_rows]
    for j, col in enumerate(t.cols, start=1):
        if col.kind == "pct" and body:
            L = get_column_letter(j)
            ws.conditional_formatting.add(
                f"{L}{body[0]}:{L}{body[-1]}",
                DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color=BLUE_LIGHT, showValue=True))
    return row + 3


def _set_widths(ws, tables):
    widths = {}
    for t in tables:
        for j, col in enumerate(t.cols, start=1):
            w = col.width or (12 if j > 1 else 30)
            widths[j] = max(widths.get(j, 0), w)
    for j, w in widths.items():
        ws.column_dimensions[get_column_letter(j)].width = min(max(w, 11), 46)


def _summary_sheet(wb, period, sections, head, meta, subtitle, audience):
    ws = wb.active
    ws.title = "Summary"
    row = _sheet_header(ws, f"Cleo Diagnostics CSP-01-001 | Monthly trial summary | {period.label}", subtitle)
    banner = ("Sponsor version, draft for review. A person checks every figure before this goes to the sponsor."
              if audience == "sponsor" else "Internal version. Contains participant-level lists. Not for the sponsor.")
    ws.cell(row=3, column=1, value=banner).font = _font(size=9, bold=True, color="D03B3B")
    completeness = "completeness" in sections

    kpis = [
        ("Enrollment", f"Participants enrolled", head["enrolled_month"], head["enrolled_to_date"], ""),
        ("Enrollment", "Enrolled per active site", round(head["per_site_month"], 2), None,
         f"{head['per_site_week']:.2f} per active site per week; {head['active_sites']} active sites"),
        ("Data completeness", "CRF forms complete", head["crf_pct"], None, f"{head['crf_done']:,} of {head['crf_expected']:,} expected forms"),
        ("Data completeness", "ePRO forms completed", head["epro_pct"], None, f"{head['epro_done']} of {head['epro_assigned']} assigned; {head['epro_open']} open"),
        ("Post-surgical", "Surgeries complete", head["surgery_month"], head["surgery_to_date"], ""),
        ("Post-surgical", "Pathology reports uploaded", head["reports_month"], head["reports_to_date"],
         f"{head['reports_pct']:.0%} of surgeries complete; median {head['turn_median']:.0f} days from enrollment"),
        ("Blood samples", "Samples registered", head["samples_month"], None, ""),
        ("Blood samples", "Samples inside registration window", head["window_pct_month"], head["window_pct_to_date"],
         f"{head['window_n_month']} of {head['window_d_month']} in {period.short}"),
        ("Data queries", "Queries raised", head["q_raised_month"], head["q_total"], f"{head['q_closed_month']} closed in {period.short}"),
        ("Data queries", "Unresolved queries at month end", None, head["q_open"] + head["q_reviewing"],
         f"{head['q_open']} open, {head['q_reviewing']} reviewing; median {head['q_median_days']:.0f} days open"),
    ]
    if not completeness:
        kpis = [k for k in kpis if k[0] != "Data completeness"]
    ws.cell(row=row, column=1, value="Key figures").font = _font(size=12, bold=True, color=NAVY)
    row += 1
    for j, h in enumerate(["Area", "Measure", period.label, "To date", "Detail"], start=1):
        c = ws.cell(row=row, column=j, value=h)
        c.font = _font(size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=BLUE)
        c.alignment = Alignment(horizontal="left" if j in (1, 2, 5) else "center", vertical="center")
    ws.row_dimensions[row].height = 22
    for i, (area, measure, m, td, detail) in enumerate(kpis):
        row += 1
        for j, v in enumerate([area, measure, m, td, detail], start=1):
            c = ws.cell(row=row, column=j, value="–" if v is None else v)
            c.font = _font(size=10, bold=j in (3, 4), color=NAVY if j in (3, 4) else INK)
            c.border = Border(bottom=THIN)
            c.alignment = Alignment(horizontal="right" if j in (3, 4) and v is not None else "center" if v is None else "left",
                                    vertical="center", wrap_text=j == 5)
            if j in (3, 4) and isinstance(v, float) and v <= 1 and "per active site" not in measure:
                c.number_format = "0%"
            elif j in (3, 4) and isinstance(v, int):
                c.number_format = "#,##0"
            if i % 2 == 1:
                c.fill = PatternFill("solid", fgColor=ZEBRA)
        ws.row_dimensions[row].height = 18
    row += 3
    for title, items in [("Highlights", narrative.highlights(head, period, completeness, audience)),
                         ("Needs attention", narrative.actions(head, sections, period, audience))]:
        ws.cell(row=row, column=1, value=title).font = _font(size=12, bold=True, color=NAVY)
        for it in items:
            row += 1
            c = ws.cell(row=row, column=1, value=f"•  {it}")
            c.font = _font(size=10)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
            ws.row_dimensions[row].height = 28
        row += 3
    for col, w in zip("ABCDE", (20, 34, 14, 12, 62)):
        ws.column_dimensions[col].width = w


def write_workbook(path: Path, period, sections, head, meta, audience="internal", chart_dir: Path | None = None):
    specs = charts.chart_specs(sections, period, head)
    images = charts.render_all(specs, chart_dir or path.parent / "charts")
    all_sections = sections
    sections = sections_for(sections, audience)
    subtitle = (f"Reporting month {period.label} · Data to {period.last_day:%d %b %Y} · "
                f"Citrus export {meta['citrus_export']:%d %b %Y %H:%M} UTC · Inventory report {meta['inventory_date']}")
    wb = Workbook()
    _summary_sheet(wb, period, all_sections if audience == "internal" else sections, head, meta, subtitle, audience)
    for key, s in sections.items():
        ws = wb.create_sheet(TAB_NAMES[key])
        row = _sheet_header(ws, s.title, subtitle)
        ws.cell(row=3, column=1, value=s.intro).font = _font(size=9.5, color=INK_2)
        shown = tables_for(s, audience)
        _set_widths(ws, shown)
        for t in shown:
            for k, sp in specs.items():
                if sp.get("before") == t.key and k in images:
                    row = _image(ws, images[k], row)
            row = _write_table(ws, t, row)
            for k, sp in specs.items():
                if sp.get("after") == t.key and k in images:
                    row = _image(ws, images[k], row)
    ws = wb.create_sheet("Definitions")
    _sheet_header(ws, "Definitions and sources", subtitle)
    r = 4
    for term, text in narrative.definitions(period, meta, "completeness" in sections):
        ws.cell(row=r, column=1, value=term).font = _font(bold=True)
        c = ws.cell(row=r, column=2, value=text)
        c.font = _font()
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top")
        ws.row_dimensions[r].height = 15 * (len(text) // 95 + 1)
        r += 1
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 28, 100
    for w in wb.worksheets:
        w.sheet_view.zoomScale = 100
        w.page_setup.orientation = "landscape"
        w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
        w.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)
    return path


def _image(ws, path, row, width=760):
    """Place a story chart below the current row and return the next free row."""
    img = XLImage(str(path))
    ratio = width / img.width
    img.width, img.height = width, int(img.height * ratio)
    ws.add_image(img, f"A{row}")
    return row + int(img.height / 20) + 3
